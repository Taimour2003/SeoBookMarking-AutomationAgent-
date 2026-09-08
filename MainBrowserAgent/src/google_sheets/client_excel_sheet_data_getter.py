# from tkinter import filedialog
# import tkinter as tk

# import pandas as pd


# def get_data_from_uploaded_file() -> list[list[str]]:
#     print("Enter in the get_data_from_uploaded_file function")
#     root = tk.Tk()
#     root.withdraw()  # Hide the main window
#     root.attributes("-topmost", True)  # Bring the file dialog to the front
#     try:
#         print("Waiting for user to select a file...")
#         file_path = filedialog.askopenfilename(
#             title="Select Excel or CSV sheet file",
#             filetypes=[("Excel/CSV Files", "*.xlsx *.xls *.csv"), ("All files", "*.*")],
#         )
#     finally:
#         root.destroy()  # Close the Tkinter root window after file selection
#     if not file_path:
#         raise RuntimeError("No file selected. Please select a valid Excel or CSV file.")

#     print(f"File selected: {file_path}")

#     if file_path.endswith(".csv"):
#         df = pd.read_csv(file_path, header=None)
#     else:
#         df = pd.read_excel(file_path, header=None)

#     df = df.fillna("")  # Replace NaN with empty strings

#     return df.astype(
#         str
#     ).values.tolist()  # Convert DataFrame to list of lists of strings


import csv
from pathlib import Path
from tkinter import filedialog
import tkinter as tk

from openpyxl import load_workbook


def get_data_from_uploaded_file() -> list[list[str]]:
    print("Enter in the get_data_from_uploaded_file function")

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    try:
        print("Waiting for user to select a file...")

        file_path = filedialog.askopenfilename(
            title="Select Excel or CSV sheet file",
            filetypes=[
                ("Excel/CSV Files", "*.xlsx *.csv"),
                ("All files", "*.*"),
            ],
        )

    finally:
        root.destroy()

    if not file_path:
        raise RuntimeError("No file selected. Please select a valid Excel or CSV file.")

    print(f"File selected: {file_path}")

    suffix = Path(file_path).suffix.lower()

    if suffix == ".csv":
        with open(
            file_path,
            "r",
            encoding="utf-8-sig",
            newline="",
        ) as file:
            return [
                ["" if cell is None else str(cell) for cell in row]
                for row in csv.reader(file)
            ]

    if suffix == ".xlsx":
        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        worksheet = workbook.active

        rows: list[list[str]] = []

        for row in worksheet.iter_rows(values_only=True):
            rows.append(["" if cell is None else str(cell) for cell in row])

        workbook.close()

        return rows

    raise RuntimeError(f"Unsupported file type: {suffix}")
